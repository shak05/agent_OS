"""Stage 7 tests: Ollama client, LLM-powered ResearchAgent and its integration.

No test here talks to a real model: client tests use a tiny fake Ollama HTTP
server on 127.0.0.1, everything else injects MockLLMClient.
"""
from __future__ import annotations

import ast
import inspect
import json
import socket
import threading
import time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation

import demo.benchmark
import demo.workflow
from agents import DocGenAgent, ResearchAgent
from agents.research import NO_SOURCES_NOTE, SYSTEM_PROMPT, build_prompt, parse_analysis
from dashboard.jobs import JobManager
from demo.benchmark import run_benchmark
from demo.fleet import register_demo_fleet
from demo.workflow import run_demo
from llm import (
    MockLLMClient,
    OllamaClient,
    OllamaConfig,
    OllamaConnectionError,
    OllamaHTTPError,
    OllamaModelNotFoundError,
    OllamaResponseError,
    OllamaTimeoutError,
)
from runtime import AgentOS

ANSWER = """SUMMARY:
Distinctive summary ALPHA-42 about the topic.

KEY FINDINGS:
- Finding ALPHA-1 from the model.
- Finding ALPHA-2 from the model.

REASONING:
- Reasoning ALPHA-R.

CONCLUSIONS:
- Conclusion ALPHA-C.

ASSUMPTIONS AND UNCERTAINTIES:
- Assumption ALPHA-U."""


# ------------------------------------------------------------ fake Ollama server
class FakeOllama:
    """Configurable stand-in for Ollama's HTTP API."""

    def __init__(self) -> None:
        self.status = 200
        self.body: Any = {"model": "llama3:latest", "message": {"role": "assistant", "content": "Hi"},
                          "done": True, "eval_count": 3}
        self.delay = 0.0
        self.requests: list[tuple[str, str, Any]] = []
        fake = self

        class Handler(BaseHTTPRequestHandler):
            def _reply(self) -> None:
                length = int(self.headers.get("Content-Length") or 0)
                raw = self.rfile.read(length) if length else b""
                fake.requests.append((self.command, self.path, json.loads(raw) if raw else None))
                time.sleep(fake.delay)
                body = fake.body if isinstance(fake.body, bytes) else json.dumps(fake.body).encode()
                self.send_response(fake.status)
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                try:
                    self.wfile.write(body)
                except OSError:
                    pass  # client gave up (timeout tests)

            do_GET = do_POST = _reply

            def log_message(self, *args: Any) -> None:
                pass

        self.server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        self.server.daemon_threads = True
        self.url = f"http://127.0.0.1:{self.server.server_address[1]}"
        threading.Thread(target=self.server.serve_forever, daemon=True).start()

    def client(self, timeout: float = 5.0, model: str = "llama3:latest") -> OllamaClient:
        return OllamaClient(OllamaConfig(base_url=self.url, model=model, timeout=timeout))


@pytest.fixture
def fake():
    server = FakeOllama()
    yield server
    server.server.shutdown()
    server.server.server_close()


MSG = [{"role": "user", "content": "hello"}]


# ---------------------------------------------------------------- client
@pytest.mark.real_http
def test_client_success_and_request_shape(fake) -> None:
    fake.body["message"]["content"] = "  Generated answer.  "
    reply = fake.client().chat(MSG)
    assert reply.content == "Generated answer."
    assert (reply.model, reply.backend, reply.eval_count) == ("llama3:latest", "ollama", 3)
    method, path, payload = fake.requests[-1]
    assert (method, path) == ("POST", "/api/chat")
    assert payload == {"model": "llama3:latest", "messages": MSG, "stream": False}


@pytest.mark.real_http
def test_client_connection_error() -> None:
    sock = socket.socket()
    sock.bind(("127.0.0.1", 0))
    port = sock.getsockname()[1]
    sock.close()  # nothing listens here now
    client = OllamaClient(OllamaConfig(base_url=f"http://127.0.0.1:{port}", timeout=5))
    with pytest.raises(OllamaConnectionError, match="Is Ollama running"):
        client.chat(MSG)


@pytest.mark.real_http
def test_client_timeout(fake) -> None:
    fake.delay = 1.0
    with pytest.raises(OllamaTimeoutError, match="did not answer within 0.2 s"):
        fake.client(timeout=0.2).chat(MSG)


@pytest.mark.real_http
def test_client_http_error(fake) -> None:
    fake.status, fake.body = 500, {"error": "out of memory"}
    with pytest.raises(OllamaHTTPError, match="HTTP 500: out of memory") as exc:
        fake.client().chat(MSG)
    assert exc.value.status == 500 and not isinstance(exc.value, OllamaModelNotFoundError)


@pytest.mark.real_http
def test_client_missing_model(fake) -> None:
    fake.status, fake.body = 404, {"error": "model 'nope:latest' not found"}
    with pytest.raises(OllamaModelNotFoundError, match="ollama pull nope:latest"):
        fake.client(model="nope:latest").chat(MSG)


@pytest.mark.real_http
@pytest.mark.parametrize("body, match", [
    (b"<html>not json</html>", "non-JSON"),
    ({"model": "llama3:latest", "done": True}, "Malformed"),
    ({"message": {"role": "assistant"}}, "Malformed"),
    ({"message": {"role": "assistant", "content": "   "}}, "empty response"),
    ({"error": "something broke"}, "something broke"),
    ([1, 2, 3], "unexpected JSON"),
])
def test_client_malformed_or_empty(fake, body, match) -> None:
    fake.body = body
    with pytest.raises(OllamaResponseError, match=match):
        fake.client().chat(MSG)


@pytest.mark.real_http
def test_client_health(fake) -> None:
    fake.body = {"models": [{"name": "llama3:latest"}]}
    assert fake.client().health() == (True, f"Ollama llama3:latest at {fake.url} is reachable")
    ok, msg = fake.client(model="mistral:latest").health()
    assert not ok and "ollama pull mistral:latest" in msg


def test_client_health_unreachable_never_raises() -> None:
    ok, msg = OllamaClient().health()  # blocked by the test guard
    assert not ok and "blocked in tests" in msg


def test_config_from_env(monkeypatch) -> None:
    for name in ("OLLAMA_BASE_URL", "OLLAMA_MODEL", "OLLAMA_TIMEOUT"):
        monkeypatch.delenv(name, raising=False)
    assert OllamaConfig.from_env() == OllamaConfig("http://localhost:11434", "llama3:latest", 300.0)
    monkeypatch.setenv("OLLAMA_BASE_URL", "http://gpu-box:11434/")
    monkeypatch.setenv("OLLAMA_MODEL", "llama3:8b")
    monkeypatch.setenv("OLLAMA_TIMEOUT", "42")
    assert OllamaConfig.from_env() == OllamaConfig("http://gpu-box:11434", "llama3:8b", 42.0)
    monkeypatch.setenv("OLLAMA_TIMEOUT", "soon")
    with pytest.raises(ValueError, match="OLLAMA_TIMEOUT"):
        OllamaConfig.from_env()


# ------------------------------------------------------------ prompt & parsing
def test_prompt_is_honest_about_no_web_access() -> None:
    assert "NO internet access" in SYSTEM_PROMPT and "independently verified" in SYSTEM_PROMPT
    plain = build_prompt("Quantum batteries")
    assert "Topic: Quantum batteries" in plain and "No source material was supplied" in plain
    with_material = build_prompt("X", question="Why?", material=["fact one", "fact two"])
    assert "Question: Why?" in with_material and "fact one\n\nfact two" in with_material


def test_parse_structured_answer() -> None:
    parsed = parse_analysis(ANSWER)
    assert parsed["summary"] == "Distinctive summary ALPHA-42 about the topic."
    assert parsed["key_points"] == ["Finding ALPHA-1 from the model.", "Finding ALPHA-2 from the model."]
    assert parsed["reasoning"] == ["Reasoning ALPHA-R."]
    assert parsed["conclusions"] == ["Conclusion ALPHA-C."]
    assert parsed["assumptions"] == ["Assumption ALPHA-U."]


def test_parse_markdown_and_inline_headings() -> None:
    text = "**Summary:** Inline summary.\n\n## Key Findings\n* one\n1. two\n\n**Conclusions**\n- done"
    parsed = parse_analysis(text)
    assert parsed["summary"] == "Inline summary."
    assert parsed["key_points"] == ["one", "two"] and parsed["conclusions"] == ["done"]


def test_parse_unstructured_answer_keeps_text() -> None:
    parsed = parse_analysis("A plain paragraph.\n\nSecond point.\n\nThird point.")
    assert parsed["summary"] == "A plain paragraph."
    assert parsed["key_points"] == ["Second point.", "Third point."]


# ------------------------------------------------------------ ResearchAgent
def test_research_agent_returns_model_output_through_agentos() -> None:
    aos = AgentOS()
    llm = MockLLMClient(response=ANSWER, model="llama3:latest")
    aos.register(ResearchAgent("research-01", agent_id="research-01", llm=llm))
    acb = aos.run_task(None, {"topic": "Edge AI"}, required_capabilities=["web_research"])
    assert acb.outcome == "completed"
    r = acb.result
    assert r["summary"] == "Distinctive summary ALPHA-42 about the topic."
    assert r["key_points"][0] == "Finding ALPHA-1 from the model."
    assert (r["model"], r["backend"], r["raw_text"]) == ("llama3:latest", "mock", ANSWER)
    assert r["sources"] == [NO_SOURCES_NOTE]  # never invents URLs
    assert "Topic: Edge AI" in llm.calls[0][-1]["content"]
    # audit: real permission name, model and backend on the completion event
    allowed = aos.audit.events("permission_allowed")
    assert [e.details["action"] for e in allowed] == ["llm_inference"]
    done = aos.audit.events("completed")[0]
    assert done.details["llm_model"] == "llama3:latest" and done.details["llm_backend"] == "mock"
    assert done.agent_id == "research-01"


def test_research_agent_lists_supplied_material_as_sources() -> None:
    aos = AgentOS()
    aos.register(ResearchAgent("r", agent_id="r", llm=MockLLMClient(response=ANSWER)))
    acb = aos.run_task("r", {"topic": "T", "material": ["doc A", "doc B"]})
    assert acb.result["sources"] == ["Supplied material #1", "Supplied material #2"]


def test_research_permission_is_not_granted_by_capability() -> None:
    aos = AgentOS()
    agent = aos.register(ResearchAgent("r", agent_id="r", llm=MockLLMClient(response=ANSWER)))
    aos.permissions.revoke("r", "llm_inference")
    assert "web_research" in agent.capabilities
    acb = aos.run_task(None, {"topic": "T"}, required_capabilities=["web_research"])
    assert acb.outcome == "denied" and agent.llm.calls == []  # model never called


def test_model_error_retries_then_succeeds() -> None:
    aos = AgentOS(max_retries=2)
    llm = MockLLMClient(response=ANSWER, error=OllamaTimeoutError("slow"), fail_times=1)
    register_demo_fleet(aos, llm=llm)
    acb = aos.run_task(None, {"topic": "T"}, required_capabilities=["web_research"])
    assert acb.outcome == "completed" and acb.agents_tried == ["research-01"] and acb.attempts == 2
    failed = aos.audit.events("failed")[0]
    assert "OllamaTimeoutError" in failed.details["error"] and failed.details["llm_backend"] == "mock"


def test_model_down_reassigns_then_escalates_without_loops() -> None:
    aos = AgentOS(max_retries=1)
    llm = MockLLMClient(error=OllamaConnectionError("Cannot reach Ollama"))
    register_demo_fleet(aos, llm=llm)
    acb = aos.run_task(None, {"topic": "T"}, required_capabilities=["web_research"])
    assert acb.outcome == "escalated"
    assert acb.agents_tried == ["research-01", "research-02", "research-03"]  # only compatible instances
    assert acb.attempts == 6 and len(llm.calls) == 6
    assert aos.audit.count("reassigned") == 2 and aos.audit.count("escalated") == 1
    assert aos.audit.count("completed") == 0  # no fabricated success


# ------------------------------------------------------------ workflow / DocGen
def _deck_text(path: Path) -> str:
    prs = Presentation(path)
    return "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)


def test_research_output_flows_into_presentation(tmp_path) -> None:
    llm = MockLLMClient(response=ANSWER, model="llama3:latest")
    aos, run = run_demo(sandbox_root=tmp_path, verbose=False, llm=llm)
    assert run.succeeded
    assert run.acbs["docgen"].task_input["research"] is run.acbs["research"].result
    text = _deck_text(tmp_path / "reports" / "ai-agents-in-healthcare.pptx")
    for marker in ("ALPHA-42", "ALPHA-1", "ALPHA-2", "ALPHA-R", "ALPHA-C", "ALPHA-U", NO_SOURCES_NOTE,
                   "mock model llama3:latest"):
        assert marker in text
    titles = [s.shapes.title.text for s in Presentation(tmp_path / "reports" / "ai-agents-in-healthcare.pptx").slides]
    assert titles == ["AI Agents in Healthcare", "Summary", "Key findings", "Supporting reasoning",
                      "Conclusions", "Assumptions & uncertainties", "Sources & method"]
    sent = json.loads((tmp_path / "sent_emails.json").read_text(encoding="utf-8"))
    assert sent[-1]["body"] == "Distinctive summary ALPHA-42 about the topic."


def test_failed_research_blocks_downstream_steps(tmp_path) -> None:
    llm = MockLLMClient(error=OllamaConnectionError("Cannot reach Ollama"))
    aos, run = run_demo(sandbox_root=tmp_path, verbose=False, llm=llm)
    assert not run.succeeded
    assert run.acbs["research"].outcome == "escalated"
    assert run.skipped == ["docgen", "file", "email"]
    assert not list(tmp_path.rglob("*.pptx")) and not (tmp_path / "sent_emails.json").exists()
    assert aos.audit.count("permission_confirmed") == 0  # email gate never reached


def test_docgen_refuses_missing_or_empty_research(tmp_path) -> None:
    aos = AgentOS(sandbox_root=tmp_path, max_retries=0)
    aos.register(DocGenAgent("docgen-01", agent_id="docgen-01"))
    for bad in [{}, {"research": None}, {"research": {"topic": "T", "summary": " ", "key_points": []}}]:
        acb = aos.run_task(None, bad, required_capabilities=["document_generation"])
        assert acb.outcome == "escalated" and "summary and key findings" in acb.error
    assert not list(tmp_path.rglob("*.pptx"))


def test_email_approval_still_required(tmp_path) -> None:
    _, run = run_demo(confirm="auto-deny", sandbox_root=tmp_path, verbose=False, llm=MockLLMClient())
    assert run.acbs["research"].outcome == "completed"
    assert run.acbs["email"].outcome == "rejected"
    assert not (tmp_path / "sent_emails.json").exists()


def test_fleet_shares_one_client_and_keeps_registry_rules() -> None:
    aos = AgentOS()
    llm = MockLLMClient()
    fleet = register_demo_fleet(aos, llm=llm)
    research = fleet["research"]
    assert [a.id for a in research] == ["research-01", "research-02", "research-03"]
    assert all(a.llm is llm for a in research)
    assert len(aos.agents) == 9
    assert [a.id for a in aos.registry.find(["web_research"])] == ["research-01", "research-02", "research-03"]


def test_default_fleet_uses_real_ollama_client(monkeypatch) -> None:
    monkeypatch.setenv("OLLAMA_MODEL", "llama3:latest")
    fleet = register_demo_fleet(AgentOS())
    client = fleet["research"][0].llm
    assert isinstance(client, OllamaClient) and client.backend == "ollama"
    assert all(a.llm is client for a in fleet["research"])


# ------------------------------------------------------------ dashboard / CLI / benchmark
def test_dashboard_job_reports_model_and_backend(tmp_path) -> None:
    mgr = JobManager(sandbox_root=tmp_path, approval_timeout=10, llm=MockLLMClient(response=ANSWER))
    assert mgr.llm.backend == "mock"
    job = mgr.start_workflow("none")
    deadline = time.time() + 10
    while job.broker.pending() is None and time.time() < deadline:
        time.sleep(0.01)
    job.broker.resolve(job.broker.pending().request_id, True)
    assert job.wait(10) and job.succeeded
    research = job.step_acb(job.steps[0])
    assert research.metadata == {"llm_model": "mock-llm", "llm_backend": "mock"}
    assert job.llm is mgr.llm


def test_dashboard_defaults_to_real_ollama() -> None:
    assert JobManager().llm.backend == "ollama"


def test_workflow_cli_mock_flag_is_explicit(monkeypatch) -> None:
    seen: dict[str, Any] = {}
    monkeypatch.setattr(demo.workflow, "run_demo", lambda **kw: seen.update(kw))
    demo.workflow.main([])
    assert seen["llm"] is None  # default: real Ollama
    demo.workflow.main(["--llm", "mock"])
    assert isinstance(seen["llm"], MockLLMClient)


def test_benchmark_never_touches_llm(monkeypatch) -> None:
    imports = {n.module for n in ast.walk(ast.parse(inspect.getsource(demo.benchmark)))
               if isinstance(n, ast.ImportFrom)}
    assert not any(m and m.startswith("llm") for m in imports)

    def forbidden(*a: Any, **k: Any) -> None:
        raise AssertionError("benchmark must not call an LLM")

    monkeypatch.setattr(OllamaClient, "chat", forbidden)
    monkeypatch.setattr(MockLLMClient, "chat", forbidden)
    results = run_benchmark(n=20, workers=4, arrival_rate=300, progress=None)
    assert [r.policy for r in results] == ["FIFO", "Priority", "Aging"]
